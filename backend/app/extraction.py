"""Bounded extraction with source locations; never creates accounting entries."""

import csv
import io
from pathlib import Path
from zipfile import BadZipFile, ZipFile

MAX_CHARACTERS = 300_000
MAX_ROWS = 5_000
MAX_PAGES = 200
EXTENSIONS = {".csv", ".xlsx", ".xls", ".pdf", ".docx", ".jpg", ".jpeg", ".png"}


class ReviewRequired(Exception):
    pass


def check_container(path: Path, extension: str):
    prefix = path.read_bytes()[:8]
    signatures = {
        ".pdf": b"%PDF-",
        ".png": b"\x89PNG\r\n\x1a\n",
        ".jpg": b"\xff\xd8\xff",
        ".jpeg": b"\xff\xd8\xff",
        ".xls": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1",
    }
    if extension in signatures and not prefix.startswith(signatures[extension]):
        raise ValueError("File contents do not match the filename extension")
    if extension in {".png", ".jpg", ".jpeg"}:
        from PIL import Image, UnidentifiedImageError

        try:
            with Image.open(path) as image:
                if image.width * image.height > 25_000_000:
                    raise ValueError("Image exceeds 25 million pixels")
                image.verify()
        except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
            raise ValueError("Invalid or excessively large image") from exc
    if extension in {".xlsx", ".docx"}:
        try:
            with ZipFile(path) as archive:
                expected = "xl/workbook.xml" if extension == ".xlsx" else "word/document.xml"
                if expected not in archive.namelist():
                    raise ValueError("Invalid Office document")
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("Uncompressed document exceeds 100 MB")
        except BadZipFile as exc:
            raise ValueError("Invalid Office document") from exc
    if extension == ".csv":
        try:
            value = path.read_text(encoding="utf-8-sig")
            if "\x00" in value:
                raise ValueError("CSV must contain text")
        except UnicodeError as exc:
            raise ValueError("Please save this CSV using UTF-8 encoding") from exc


def extract(path: Path, extension: str) -> dict:
    blocks = []
    total = 0

    def add(location, text):
        nonlocal total
        if total + len(text) > MAX_CHARACTERS:
            raise ReviewRequired("Document exceeds extraction limit; split it into smaller files")
        total += len(text)
        blocks.append({"location": location, "text": text})

    if extension == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            for index, row in enumerate(csv.reader(handle), 1):
                if index > MAX_ROWS:
                    raise ReviewRequired("CSV exceeds 5,000 rows; split it into smaller files")
                add({"row": index}, _row_text(row))
    elif extension == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            count = 0
            for sheet in workbook:
                for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    count += 1
                    if count > MAX_ROWS:
                        raise ReviewRequired("Workbook exceeds 5,000 rows")
                    add({"sheet": sheet.title, "row": index}, _row_text(row))
        finally:
            workbook.close()
    elif extension == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(path, on_demand=True)
        try:
            count = 0
            for sheet in workbook.sheets():
                for index in range(sheet.nrows):
                    count += 1
                    if count > MAX_ROWS:
                        raise ReviewRequired("Workbook exceeds 5,000 rows")
                    add({"sheet": sheet.name, "row": index + 1}, _row_text(sheet.row_values(index)))
        finally:
            workbook.release_resources()
    elif extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ReviewRequired("Unlock this PDF before uploading it")
        if len(reader.pages) > MAX_PAGES:
            raise ReviewRequired("PDF exceeds 200 pages; split it into smaller files")
        for index, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                raise ReviewRequired(
                    f"Page {index} requires OCR; OCR is not enabled in this release"
                )
            add({"page": index}, text)
    elif extension == ".docx":
        from docx import Document

        document = Document(path)
        for index, paragraph in enumerate(document.paragraphs, 1):
            if paragraph.text.strip():
                add({"paragraph": index}, paragraph.text)
        for table_index, table in enumerate(document.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                add({"table": table_index, "row": row_index}, _row_text(c.text for c in row.cells))
    else:
        raise ReviewRequired("This image requires OCR; OCR is not enabled in this release")
    if not any(block["text"].strip() for block in blocks):
        raise ReviewRequired("No readable content found")
    return {
        "blocks": blocks,
        "verified": False,
        "note": "Extracted text is not verified accounting data",
    }


def _row_text(row):
    output = io.StringIO()
    csv.writer(output).writerow("" if cell is None else str(cell) for cell in row)
    return output.getvalue().strip()
