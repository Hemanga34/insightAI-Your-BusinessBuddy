from io import BytesIO

from conftest import register
from docx import Document as WordDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def test_text_pdf_has_page_reference(client):
    headers = register(client)
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 200 Td (Invoice total INR 1200) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    response = client.post(
        "/api/documents", headers=headers, files={"file": ("invoice.pdf", output.getvalue())}
    )
    assert response.json()["status"] == "extracted"
    detail = client.get(f"/api/documents/{response.json()['id']}", headers=headers).json()
    block = detail["extraction"]["blocks"][0]
    assert block["location"] == {"page": 1}
    assert "INR 1200" in block["text"]


def test_docx_includes_tables_and_paragraphs(client):
    headers = register(client)
    document = WordDocument()
    document.add_paragraph("October statement")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "INR 500"
    output = BytesIO()
    document.save(output)
    response = client.post(
        "/api/documents", headers=headers, files={"file": ("statement.docx", output.getvalue())}
    )
    assert response.json()["status"] == "extracted"
    detail = client.get(f"/api/documents/{response.json()['id']}", headers=headers).json()
    blocks = detail["extraction"]["blocks"]
    assert blocks[0]["location"] == {"paragraph": 1}
    assert blocks[1]["location"] == {"table": 1, "row": 1}
    assert "500" in blocks[1]["text"]


def test_scan_pdf_and_large_upload_return_explicit_states(client):
    headers = register(client)
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = BytesIO()
    writer.write(output)
    response = client.post(
        "/api/documents", headers=headers, files={"file": ("scan.pdf", output.getvalue())}
    )
    assert response.json()["status"] == "needs_attention"
    assert "OCR" in response.json()["error"]
    from app.config import get_settings

    get_settings().max_upload_bytes = 4
    response = client.post(
        "/api/documents", headers=headers, files={"file": ("too-large.csv", b"amount\n500")}
    )
    assert response.status_code == 413
